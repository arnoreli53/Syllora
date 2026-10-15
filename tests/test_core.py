from __future__ import annotations

import base64
import hashlib
import http.client
import http.server
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

PYTHON_DIR = Path(__file__).resolve().parents[1]
if str(PYTHON_DIR) not in sys.path:
    sys.path.insert(0, str(PYTHON_DIR))

from PySide6.QtCore import QDateTime, QEvent
from PySide6.QtSql import QSqlQuery
from PySide6.QtWidgets import QApplication

import ai_backend
import chatgpt_auth
import migrate
import paths
import release_metadata
import updater
import ui_settings
import user_profiles
from db_qt import close_db, open_db
from ui_courses import calculate_previous_course_gpa_metrics, list_current_course_completion_state
from ui_settings import (
    SET_GRADE_CALC_MODE,
    SET_GRADE_DISPLAY,
    SET_OPENAI_API_KEY,
    _import_backup_json,
    _set,
    _settings_table_snapshot,
    set_str,
)
from ui_tasks import import_tasks_from_excel, load_workbook
from version import APP_BUNDLE_IDENTIFIER, APP_BUNDLE_NAME, GITHUB_REPO_SLUG


APP = QApplication.instance() or QApplication([])


def execute(sql: str, values: tuple = ()) -> None:
    query = QSqlQuery()
    query.prepare(sql)
    for value in values:
        query.addBindValue(value)
    if not query.exec():
        raise AssertionError(query.lastError().text())


def scalar(sql: str):
    query = QSqlQuery()
    if not query.exec(sql):
        raise AssertionError(query.lastError().text())
    if not query.next():
        raise AssertionError(f"Query returned no rows: {sql}")
    return query.value(0)


class DatabaseBehaviorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        data_dir = Path(self.temp_dir.name)
        paths.set_active_data_paths(data_dir / "syllora.db", data_dir / "settings.ini")
        open_db()
        migrate.ensure_schema()

    def tearDown(self) -> None:
        close_db()
        self.temp_dir.cleanup()

    def test_database_safety_pragmas_are_enabled(self) -> None:
        self.assertEqual(int(scalar("PRAGMA foreign_keys")), 1)
        self.assertGreaterEqual(int(scalar("PRAGMA busy_timeout")), 5000)
        self.assertEqual(str(scalar("PRAGMA journal_mode")).lower(), "wal")

    def test_backup_snapshot_excludes_api_key(self) -> None:
        _set(SET_OPENAI_API_KEY, "test-secret-that-must-not-be-backed-up")
        _set("ordinary_setting", "kept")
        snapshot = _settings_table_snapshot()
        self.assertNotIn(SET_OPENAI_API_KEY, snapshot)
        self.assertEqual(snapshot["ordinary_setting"], "kept")

    def test_backup_restore_is_atomic_and_preserves_unknown_due_dates(self) -> None:
        backup_path = Path(self.temp_dir.name) / "backup.json"
        backup_path.write_text(
            json.dumps(
                {
                    "courses": [{"name": "Restored Course", "term": "fall", "academic_year_start": 2026}],
                    "tasks": [{"course": "Restored Course", "item": "Undated task", "due_datetime": ""}],
                    "previous_courses": [],
                    "settings": {
                        SET_OPENAI_API_KEY: "must-not-restore",
                        "ordinary_setting": "restored",
                    },
                    "app_qsettings": {"post_commit_preference": "restored"},
                }
            ),
            encoding="utf-8",
        )

        with patch.object(ui_settings, "_ensure_course", side_effect=RuntimeError("injected restore failure")):
            with self.assertRaisesRegex(RuntimeError, "injected restore failure"):
                _import_backup_json(str(backup_path))

        self.assertEqual(int(scalar("SELECT count(*) FROM courses")), 0)
        self.assertEqual(int(scalar("SELECT count(*) FROM settings")), 0)
        self.assertIsNone(ui_settings.app_qsettings().value("post_commit_preference"))

        _import_backup_json(str(backup_path))
        self.assertEqual(int(scalar("SELECT due_datetime IS NULL FROM tasks WHERE item='Undated task'")), 1)
        self.assertEqual(int(scalar(f"SELECT count(*) FROM settings WHERE key='{SET_OPENAI_API_KEY}'")), 0)
        self.assertEqual(scalar("SELECT value FROM settings WHERE key='ordinary_setting'"), "restored")
        self.assertEqual(ui_settings.app_qsettings().value("post_commit_preference"), "restored")

    def test_submitted_without_grade_can_count_as_zero(self) -> None:
        execute(
            "INSERT INTO courses(name, term, academic_year_start) VALUES(?, ?, ?)",
            ("Test Course", "fall", 2026),
        )
        course_id = int(scalar("SELECT id FROM courses WHERE name='Test Course'"))
        execute(
            "INSERT INTO tasks(course_id, item, status, weight, grade) VALUES(?, ?, ?, ?, ?)",
            (course_id, "Graded", "graded", 50.0, 100.0),
        )
        execute(
            "INSERT INTO tasks(course_id, item, status, weight, grade) VALUES(?, ?, ?, ?, ?)",
            (course_id, "Submitted", "submitted", 50.0, None),
        )

        set_str(SET_GRADE_CALC_MODE, "graded_only")
        self.assertAlmostEqual(list_current_course_completion_state()[course_id]["final_percent"], 100.0)

        set_str(SET_GRADE_CALC_MODE, "submitted_as_zero")
        self.assertAlmostEqual(list_current_course_completion_state()[course_id]["final_percent"], 50.0)

    def test_previous_course_gpa_uses_selected_display_scale(self) -> None:
        rows = [
            {
                "id": 1,
                "course_name": "A Plus Course",
                "final_percent": 95.0,
                "final_letter": "",
                "credit_hours": 3.0,
                "term": "fall",
                "academic_year_start": 2025,
                "created_at": "2025-09-01",
            },
            {
                "id": 2,
                "course_name": "A Course",
                "final_percent": -1.0,
                "final_letter": "A",
                "credit_hours": 3.0,
                "term": "winter",
                "academic_year_start": 2025,
                "created_at": "2026-01-01",
            },
        ]

        set_str(SET_GRADE_DISPLAY, "gpa4.0")
        metrics_40 = calculate_previous_course_gpa_metrics(rows, 6)
        self.assertAlmostEqual(metrics_40["cumulative"]["gpa"], 4.0)
        self.assertAlmostEqual(metrics_40["last_x"]["gpa"], 4.0)
        self.assertEqual(metrics_40["cumulative"]["scale"], "4.0")

        set_str(SET_GRADE_DISPLAY, "gpa4.3")
        metrics_43 = calculate_previous_course_gpa_metrics(rows, 6)
        self.assertAlmostEqual(metrics_43["cumulative"]["gpa"], 4.15)
        self.assertAlmostEqual(metrics_43["last_x"]["gpa"], 4.15)
        self.assertEqual(metrics_43["cumulative"]["scale"], "4.3")

    @unittest.skipIf(load_workbook is None, "openpyxl is not installed")
    def test_excel_import_keeps_missing_due_date_unknown(self) -> None:
        from openpyxl import Workbook

        workbook_path = Path(self.temp_dir.name) / "tasks.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Tasks"
        sheet.append(["Course", "Item", "Due Date"])
        sheet.append(["Undated Course", "Research topic", None])
        workbook.save(workbook_path)

        _, tasks_added = import_tasks_from_excel(str(workbook_path))
        self.assertEqual(tasks_added, 1)
        self.assertEqual(int(scalar("SELECT due_datetime IS NULL FROM tasks WHERE item='Research topic'")), 1)
        self.assertEqual(scalar("SELECT due_status FROM tasks WHERE item='Research topic'"), "unknown")

    def test_schema_creation_rolls_back_on_failure(self) -> None:
        close_db()
        data_dir = Path(self.temp_dir.name)
        failing_db = data_dir / "failing.db"
        paths.set_active_data_paths(failing_db, data_dir / "failing.ini")
        open_db()

        original_exec = migrate._exec
        call_count = 0

        def fail_after_first_statement(sql: str) -> None:
            nonlocal call_count
            original_exec(sql)
            call_count += 1
            if call_count == 2:
                raise RuntimeError("injected migration failure")

        with patch.object(migrate, "_exec", side_effect=fail_after_first_statement):
            with self.assertRaisesRegex(RuntimeError, "injected migration failure"):
                migrate.ensure_schema()

        self.assertEqual(int(scalar("SELECT count(*) FROM sqlite_master WHERE type='table' AND name IN ('courses','previous_courses')")), 0)

    def test_legacy_task_rebuild_keeps_rows_with_foreign_keys_enabled(self) -> None:
        close_db()
        data_dir = Path(self.temp_dir.name)
        legacy_path = data_dir / "legacy.db"
        with sqlite3.connect(legacy_path) as connection:
            connection.executescript(
                """
                CREATE TABLE courses (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE);
                CREATE TABLE tasks (
                    id INTEGER PRIMARY KEY,
                    course_id INTEGER NOT NULL,
                    project_id INTEGER,
                    item TEXT NOT NULL,
                    component TEXT DEFAULT '',
                    due_datetime TEXT,
                    status TEXT NOT NULL DEFAULT 'not started',
                    weight REAL,
                    grade REAL,
                    ungraded INTEGER NOT NULL DEFAULT 0,
                    notes TEXT DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(course_id) REFERENCES courses(id) ON DELETE CASCADE
                );
                INSERT INTO courses(id, name) VALUES(1, 'Legacy Course');
                INSERT INTO tasks(id, course_id, project_id, item) VALUES(1, 1, 99, 'Legacy Task');
                """
            )

        paths.set_active_data_paths(legacy_path, data_dir / "legacy.ini")
        open_db()
        migrate.ensure_schema()

        columns = {str(row[1]) for row in sqlite3.connect(legacy_path).execute("PRAGMA table_info(tasks)")}
        self.assertNotIn("project_id", columns)
        self.assertIn("priority", columns)
        self.assertEqual(scalar("SELECT item FROM tasks WHERE id=1"), "Legacy Task")


class UpdaterSafetyTests(unittest.TestCase):
    def test_latest_release_manifest_is_preferred(self) -> None:
        self.assertEqual(updater.LATEST_JSON_URLS[0], updater.LATEST_RELEASE_MANIFEST_URL)
        self.assertEqual(
            updater.LATEST_RELEASE_MANIFEST_URL,
            f"https://github.com/{GITHUB_REPO_SLUG}/releases/latest/download/latest.json",
        )

    def test_manifest_requires_official_release_and_valid_checksum(self) -> None:
        base = {
            "latest_version": "99.0.0",
            "macos_zip_url": f"https://github.com/{GITHUB_REPO_SLUG}/releases/download/v99/Syllora.zip",
            "sha256": "a" * 64,
        }
        with patch.object(updater, "_load_json_from_url", return_value=dict(base)):
            self.assertTrue(updater.check_for_updates()["update_available"])

        outside = dict(base, macos_zip_url="https://example.com/Syllora.zip")
        with patch.object(updater, "_load_json_from_url", return_value=outside):
            with self.assertRaisesRegex(RuntimeError, "outside the official"):
                updater.check_for_updates()

        invalid_hash = dict(base, sha256="not-a-checksum")
        with patch.object(updater, "_load_json_from_url", return_value=invalid_hash):
            with self.assertRaisesRegex(RuntimeError, "invalid SHA-256"):
                updater.check_for_updates()

    def test_newest_valid_manifest_wins_over_stale_latest_redirect(self) -> None:
        stale = {
            "latest_version": "1.0.9.1.1",
            "macos_zip_url": (
                f"https://github.com/{GITHUB_REPO_SLUG}/releases/download/"
                "auto-v1.0.9.1.1/Syllora-1.0.9.1.1-macos.zip"
            ),
            "sha256": "a" * 64,
        }
        fresh = {
            "latest_version": "1.0.9.2.1",
            "macos_zip_url": (
                f"https://github.com/{GITHUB_REPO_SLUG}/releases/download/"
                "auto-v1.0.9.2.1/Syllora-1.0.9.2.1-macos.zip"
            ),
            "sha256": "b" * 64,
        }
        responses = [stale, fresh, fresh, fresh]
        with (
            patch.object(updater, "_load_json_from_url", side_effect=responses),
            patch.object(updater, "current_update_version", return_value="1.0.9.1.1"),
        ):
            result = updater.check_for_updates()

        self.assertTrue(result["update_available"])
        self.assertEqual(result["latest_version"], "1.0.9.2.1")
        self.assertEqual(result["sha256"], "b" * 64)

    def test_private_or_missing_manifest_has_actionable_error(self) -> None:
        not_found = urllib.error.HTTPError(
            updater.LATEST_RELEASE_MANIFEST_URL,
            404,
            "Not Found",
            None,
            None,
        )
        with patch.object(updater, "_load_json_from_url", side_effect=not_found):
            with self.assertRaisesRegex(RuntimeError, "repository may be private"):
                updater.check_for_updates()

    def test_installer_checks_bundle_identity_and_signature(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            script_path = updater.write_update_installer_script(
                root / "Syllora.zip",
                root / APP_BUNDLE_NAME,
            )
            script = script_path.read_text(encoding="utf-8")
            self.assertIn(APP_BUNDLE_IDENTIFIER, script)
            self.assertIn("codesign --verify --deep --strict", script)
            self.assertIn('if ! wait_for_pid "$APP_PID"', script)
            self.assertIn('/bin/rm -rf "$RUNTIME_DIR"', script)
            syntax_check = subprocess.run(
                ["/bin/zsh", "-n", str(script_path)],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(syntax_check.returncode, 0, syntax_check.stderr)


class ReleaseMetadataTests(unittest.TestCase):
    def test_automated_release_manifest_matches_github_asset(self) -> None:
        manifest = release_metadata.build_manifest(
            version="1.0.9.42.1",
            repository=GITHUB_REPO_SLUG,
            release_tag="auto-v1.0.9.42.1",
            asset_name="Syllora-1.0.9.42.1-macos.zip",
            sha256="A" * 64,
            release_notes=["Automated test release."],
        )
        self.assertEqual(manifest["latest_version"], "1.0.9.42.1")
        self.assertEqual(manifest["sha256"], "a" * 64)
        self.assertEqual(
            manifest["macos_zip_url"],
            "https://github.com/arnoreli53/Syllora/releases/download/"
            "auto-v1.0.9.42.1/Syllora-1.0.9.42.1-macos.zip",
        )

    def test_release_metadata_rejects_unsafe_values(self) -> None:
        common = {
            "version": "1.0.9",
            "repository": GITHUB_REPO_SLUG,
            "release_tag": "v1.0.9",
            "asset_name": "Syllora-1.0.9-macos.zip",
            "sha256": "a" * 64,
        }
        with self.assertRaises(ValueError):
            release_metadata.build_manifest(**dict(common, release_tag="../unsafe"))
        with self.assertRaises(ValueError):
            release_metadata.build_manifest(**dict(common, sha256="invalid"))


class UiSmokeTests(unittest.TestCase):
    def test_main_window_builds_at_a_laptop_friendly_size(self) -> None:
        import ui_main
        import ui_settings

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths.set_active_data_paths(root / "ui.db", root / "settings.ini")
            open_db()
            migrate.ensure_schema()
            try:
                with (
                    patch.object(ui_settings, "current_profile", return_value=None),
                    patch.object(ui_settings, "list_profiles", return_value=[]),
                    patch.object(ui_settings, "load_secret", return_value=""),
                    patch.object(chatgpt_auth, "_read_blob", return_value=""),
                    patch.object(ui_main, "current_profile", return_value=None),
                    patch.object(ui_main, "list_profiles", return_value=[]),
                ):
                    window = ui_main.MainWindow()
                    window.show()
                    APP.processEvents()
                    self.assertEqual(window.tab_widget.count(), 4)
                    self.assertEqual([button.text() for button in window.nav_buttons], ["Overview", "Tasks", "Courses", "Settings"])
                    self.assertFalse(window.tab_widget.tabBar().isVisible())
                    self.assertEqual(window.settings_page._settings_category_tabs.count(), 4)
                    self.assertLessEqual(window.minimumWidth(), 820)
                    self.assertLessEqual(window.minimumHeight(), 600)

                    settings = window.settings_page
                    self.assertEqual(settings.lbl_chatgpt_status.text(), "Not signed in")
                    self.assertEqual(settings.btn_chatgpt.text(), "Continue with ChatGPT")
                    self.assertTrue(settings.lbl_chatgpt_manage.isHidden())
                    signed_in = json.dumps(
                        {
                            "client_id": "oaiapp_x",
                            "last_email": "student@example.com",
                            "session": {
                                "client_id": "oaiapp_x", "access_token": "a", "refresh_token": "r", "expires_at": time.time() + 3600,
                                "scopes": ["openid", chatgpt_auth.PLAN_USAGE_SCOPE], "subject": "u", "email": "student@example.com", "name": "",
                            },
                        }
                    )
                    with patch.object(chatgpt_auth, "_read_blob", return_value=signed_in):
                        settings._update_chatgpt_status()
                        self.assertIn("student@example.com", settings.lbl_chatgpt_status.text())
                        self.assertIn("Using ChatGPT plan", settings.lbl_chatgpt_status.text())
                        self.assertEqual(settings.btn_chatgpt.text(), "Sign out")
                        self.assertFalse(settings.lbl_chatgpt_manage.isHidden())
                    settings._update_chatgpt_status()
                    screenshot_path = os.environ.get("SYLLORA_UI_SCREENSHOT", "").strip()
                    if screenshot_path:
                        requested_tab = os.environ.get("SYLLORA_UI_TAB", "").strip().lower()
                        for index in range(window.tab_widget.count()):
                            if window.tab_widget.tabText(index).lower() == requested_tab:
                                window.tab_widget.setCurrentIndex(index)
                                APP.processEvents()
                                break
                        self.assertTrue(window.grab().save(screenshot_path))
                    window._clear_pages()
                    window.setCentralWidget(None)
                    window.close()
                    window.deleteLater()
                    del window
                    APP.sendPostedEvents(None, QEvent.Type.DeferredDelete)
                    APP.processEvents()
            finally:
                close_db()


class OverviewReminderTests(unittest.TestCase):
    def test_startup_reminder_is_held_until_app_is_live_then_delivered_once(self) -> None:
        import ui_overview

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths.set_active_data_paths(root / "reminders.db", root / "settings.ini")
            open_db()
            migrate.ensure_schema()
            try:
                execute(
                    "INSERT INTO courses(name, term, academic_year_start) VALUES(?, ?, ?)",
                    ("REM 101", "fall", 2026),
                )
                course_id = scalar("SELECT id FROM courses WHERE name = 'REM 101'")
                due = QDateTime.currentDateTime().addDays(1).toString("yyyy-MM-dd HH:mm")
                execute(
                    "INSERT INTO tasks(course_id, item, status, due_datetime) VALUES(?, ?, ?, ?)",
                    (course_id, "Essay", "not started", due),
                )

                sent: list[tuple[str, str]] = []
                with patch.object(
                    ui_overview.OverviewPage,
                    "_show_desktop_notification",
                    lambda self, title, message: sent.append((title, message)),
                ):
                    page = ui_overview.OverviewPage()
                    page.refresh()
                    APP.processEvents()
                    # The OS drops notifications posted before the app has finished launching, so the
                    # startup refreshes must not send (or mark as sent) anything yet.
                    self.assertEqual(sent, [])

                    page._reminder_startup_timer.stop()
                    page._enable_reminder_notifications()
                    self.assertEqual(len(sent), 1)
                    self.assertEqual(sent[0][0], "Course reminders")

                    page.refresh()
                    self.assertEqual(len(sent), 1, "an unchanged reminder must not be re-sent")

                    page.deleteLater()
                    APP.sendPostedEvents(None, QEvent.Type.DeferredDelete)
            finally:
                close_db()


RSA_TEST_MODULUS_HEX = "AA8EF81C287E2A5188F4864517C69666AF6599EC7A575EB12F9D8F0792831AF4C111260BFF5BE410B004D1897D8803252B7BF9CA043E4B6B62C5BC006DF6CB03E2CF71F7E6D949E441BBED2790FEFA5DE4867F2D7C8274CB338BFE80E258BF0A3E4917F00D2AFF08AD85B14AE47A00EC2D35EEB4752FAD28197CDBEBEE83F35D1E6EB8B5DE9126E1662F851397DABF36E1F24AA705758167CD07761E9B97749A47D4DB154D4438D1978B9E3EA3A2D15936E11AAD876C591981D89E985699DF8EB4849A23E82F81C3BFC45D9A4F45EDFBD3348AD692C880A015D465F5D3396E94943F898F0A19FA9613F3CAA533EC7F8C8B541C20F3F73FE1C5A73ED96250C8B3"
RSA_TEST_SIGNATURE_HEX = "5a8e857060f710e48eab18985bad1a3f2c2c8a9bd56c8cf21bd278f8f64deb5f838de11635534f17897ca7621338e3881a887468b5cd23805adce0a9daaa991bc0dc4fdbc9c6e4cdf08bac6adcf0c555e4801d55978d7b8594a2b4ea1eb919148191d2be7633bbed6010b72ba126e2aa01bf6b054f54335e484565e6e82717b30cd59d58480acb89a307f63a10fef575da6ae80dabe65e3ae25c76242e11718464b473827f1d9e3c4e8a30841ee57dbdc7ddfbbfadf2e5b03d36b58f97e224a239efc118122fe41f539f790c897d39a5515bcd3196c8f75a0eefe04439e9228fb71dec15d1ab97503ec26ddcc83568bb308911ee1db5d5ff2e62a55914295cf4"


class _MemoryKeychain:
    """Stands in for the macOS Keychain so tests never read or write the real one."""

    def __init__(self) -> None:
        self.blobs: dict[str, str] = {}

    def patches(self):
        return (
            patch.object(chatgpt_auth, "_read_blob", side_effect=lambda pid: self.blobs.get(pid, "")),
            patch.object(chatgpt_auth, "_write_blob", side_effect=lambda pid, text: self.blobs.__setitem__(pid, text)),
            patch.object(chatgpt_auth, "_delete_blob", side_effect=lambda pid: self.blobs.pop(pid, None)),
        )


class _FakeTokenServer:
    def __init__(self, responder) -> None:
        outer = self
        self.requests: list[dict[str, str]] = []

        class Handler(http.server.BaseHTTPRequestHandler):
            def log_message(self, *_args) -> None:
                return

            def do_POST(self) -> None:  # noqa: N802
                body = urllib.parse.parse_qs(self.rfile.read(int(self.headers["Content-Length"])).decode())
                outer.requests.append({k: v[0] for k, v in body.items()})
                status, payload = responder(outer.requests[-1])
                data = json.dumps(payload).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

        self._httpd = http.server.HTTPServer(("127.0.0.1", 0), Handler)
        self.url = f"http://127.0.0.1:{self._httpd.server_address[1]}/token"
        threading.Thread(target=self._httpd.serve_forever, daemon=True).start()

    def close(self) -> None:
        self._httpd.shutdown()
        self._httpd.server_close()


def _fake_jwt(claims: dict, *, alg: str = "RS256", kid: str = "k1") -> str:
    def encode(obj) -> str:
        return base64.urlsafe_b64encode(json.dumps(obj).encode()).rstrip(b"=").decode()

    return f"{encode({'alg': alg, 'kid': kid})}.{encode(claims)}.{encode('signature')}"


class ChatGPTAuthTests(unittest.TestCase):
    def setUp(self) -> None:
        self.keychain = _MemoryKeychain()
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        stack = [
            *self.keychain.patches(),
            patch.object(chatgpt_auth, "HOST_ID_PATH", Path(self._tmp.name) / "host_id"),
            patch.object(chatgpt_auth, "_fetch_jwks", return_value=[{"kid": "k1", "kty": "RSA"}]),
            patch.object(chatgpt_auth, "_verify_signature", return_value=True),
        ]
        for item in stack:
            item.start()
            self.addCleanup(item.stop)

    def _query(self, url: str) -> dict[str, str]:
        return {k: v[0] for k, v in urllib.parse.parse_qs(urllib.parse.urlparse(url).query).items()}

    def test_pkce_challenge_matches_verifier(self) -> None:
        verifier, challenge = chatgpt_auth.generate_pkce()
        digest = hashlib.sha256(verifier.encode("ascii")).digest()
        self.assertEqual(challenge, base64.urlsafe_b64encode(digest).rstrip(b"=").decode())
        self.assertNotEqual(chatgpt_auth.generate_pkce()[0], verifier)

    def test_authorize_url_registers_first_then_reuses_the_issued_client(self) -> None:
        common = dict(host_id="urn:uuid:h", redirect_uri="http://127.0.0.1:1455/auth/callback", state="s", nonce="n", code_challenge="c")
        first = self._query(chatgpt_auth.build_authorize_url(client_id=chatgpt_auth.DYNAMIC_CLIENT_ID, **common))
        self.assertEqual(first["client_id"], "dynamic_agent_client")
        self.assertEqual(first["agent_name_hint"], chatgpt_auth.APP_DISPLAY_NAME)
        self.assertEqual(first["ext_agent_host_id"], "urn:uuid:h")
        self.assertEqual(first["redirect_uri"], common["redirect_uri"])
        self.assertEqual(first["code_challenge_method"], "S256")
        self.assertEqual(first["resource"], "https://api.openai.com/v1")
        self.assertIn("chatgpt.tokens.use.direct", first["scope"].split())
        self.assertNotIn("login_hint", first)

        again = self._query(chatgpt_auth.build_authorize_url(client_id="oaiapp_x", login_hint="a@b.co", **common))
        self.assertEqual(again["client_id"], "oaiapp_x")
        self.assertNotIn("agent_name_hint", again)
        self.assertEqual(again["login_hint"], "a@b.co")

    def test_rsa_signature_verification(self) -> None:
        modulus = int(RSA_TEST_MODULUS_HEX, 16)
        signature = bytes.fromhex(RSA_TEST_SIGNATURE_HEX)
        verify = chatgpt_auth._rsa_verify_pkcs1_sha256
        self.assertTrue(verify(modulus, 65537, b"header.payload", signature))
        self.assertFalse(verify(modulus, 65537, b"header.payload2", signature))
        self.assertFalse(verify(modulus, 65537, b"header.payload", signature[:-1] + bytes([signature[-1] ^ 1])))
        self.assertFalse(verify(modulus, 65537, b"header.payload", signature[1:]))

    def test_id_token_validation_rejects_bad_tokens(self) -> None:
        good = {"iss": chatgpt_auth.AUTH_ISSUER, "aud": ["oaiapp_x"], "exp": time.time() + 600, "nonce": "n1", "sub": "u"}
        self.assertEqual(chatgpt_auth.validate_id_token(_fake_jwt(good), client_id="oaiapp_x", nonce="n1")["sub"], "u")
        for label, token, nonce in (
            ("nonce", _fake_jwt(good), "other"),
            ("audience", _fake_jwt({**good, "aud": "someone_else"}), "n1"),
            ("issuer", _fake_jwt({**good, "iss": "https://evil.example"}), "n1"),
            ("expired", _fake_jwt({**good, "exp": time.time() - 3600}), "n1"),
            ("alg none", _fake_jwt(good, alg="none"), "n1"),
            ("unknown key", _fake_jwt(good, kid="missing"), "n1"),
            ("malformed", "not-a-jwt", "n1"),
        ):
            with self.subTest(label), self.assertRaises(chatgpt_auth.ChatGPTAuthError):
                chatgpt_auth.validate_id_token(token, client_id="oaiapp_x", nonce=nonce)
        with patch.object(chatgpt_auth, "_verify_signature", return_value=False), self.assertRaises(chatgpt_auth.ChatGPTAuthError):
            chatgpt_auth.validate_id_token(_fake_jwt(good), client_id="oaiapp_x", nonce="n1")

    def test_callback_server_only_accepts_our_state_on_the_loopback_host(self) -> None:
        server = chatgpt_auth._CallbackServer("good-state", preferred_port=0)
        self.addCleanup(server.close)
        codes: list[int] = []

        def browse() -> None:
            for host, path in (
                (f"127.0.0.1:{server.port}", "/auth/callback?code=x&state=wrong"),
                (f"127.0.0.1:{server.port}", "/elsewhere?state=good-state"),
                (f"evil.example:{server.port}", "/auth/callback?code=x&state=good-state"),
                (f"127.0.0.1:{server.port}", "/auth/callback?code=abc&state=good-state&client_id=oaiapp_1"),
            ):
                conn = http.client.HTTPConnection("127.0.0.1", server.port, timeout=5)
                conn.request("GET", path, headers={"Host": host})
                codes.append(conn.getresponse().status)
                conn.close()

        thread = threading.Thread(target=browse)
        thread.start()
        result = server.wait(10, None)
        thread.join()
        self.assertEqual(codes, [400, 404, 403, 200])
        self.assertEqual((result["code"], result["client_id"]), ("abc", "oaiapp_1"))

    def test_callback_server_stops_waiting_when_cancelled(self) -> None:
        server = chatgpt_auth._CallbackServer("s", preferred_port=0)
        self.addCleanup(server.close)
        cancel = threading.Event()
        threading.Timer(0.3, cancel.set).start()
        with self.assertRaises(chatgpt_auth.ChatGPTSignInCancelled):
            server.wait(10, cancel)

    def _sign_in_against_fake_server(self, profile: str):
        captured: dict[str, str] = {}

        def respond(form: dict[str, str]):
            claims = {
                "iss": chatgpt_auth.AUTH_ISSUER, "aud": [form["client_id"]], "exp": time.time() + 3600,
                "nonce": captured["nonce"], "sub": "user-1", "email": "student@example.com", "name": "Sam Student",
            }
            return 200, {
                "access_token": "access-1", "refresh_token": "refresh-1", "id_token": _fake_jwt(claims),
                "expires_in": 3600, "scope": chatgpt_auth.SCOPES,
            }

        fake = _FakeTokenServer(respond)
        self.addCleanup(fake.close)

        def open_url(url: str) -> bool:
            query = self._query(url)
            captured.update(query)

            def callback() -> None:
                params = {"code": "auth-code", "state": query["state"]}
                if query["client_id"] == chatgpt_auth.DYNAMIC_CLIENT_ID:
                    params["client_id"] = "oaiapp_issued"
                target = urllib.parse.urlparse(query["redirect_uri"])
                conn = http.client.HTTPConnection(target.hostname, target.port, timeout=5)
                conn.request("GET", f"{target.path}?{urllib.parse.urlencode(params)}")
                conn.getresponse().read()
                conn.close()

            threading.Thread(target=callback).start()
            return True

        with patch.object(chatgpt_auth, "TOKEN_URL", fake.url):
            session = chatgpt_auth.sign_in(profile, open_url=open_url, timeout=10)
        return session, fake, captured

    def test_sign_in_registers_a_client_and_stores_the_session(self) -> None:
        session, fake, authorize = self._sign_in_against_fake_server("p1")

        self.assertEqual(authorize["client_id"], "dynamic_agent_client")
        self.assertEqual((session.client_id, session.email, session.name), ("oaiapp_issued", "student@example.com", "Sam Student"))
        self.assertTrue(session.plan_usage_enabled)
        form = fake.requests[0]
        self.assertEqual(form["grant_type"], "authorization_code")
        self.assertEqual(form["client_id"], "oaiapp_issued")
        self.assertEqual(form["redirect_uri"], authorize["redirect_uri"])
        challenge = base64.urlsafe_b64encode(hashlib.sha256(form["code_verifier"].encode()).digest()).rstrip(b"=").decode()
        self.assertEqual(challenge, authorize["code_challenge"])

        status = chatgpt_auth.status("p1")
        self.assertEqual((status.signed_in, status.email, status.plan_usage), (True, "student@example.com", True))

        # Signing out keeps the registration so the next sign-in doesn't register a new client.
        chatgpt_auth.sign_out("p1")
        self.assertFalse(chatgpt_auth.status("p1").signed_in)
        _, _, second_authorize = self._sign_in_against_fake_server("p1")
        self.assertEqual(second_authorize["client_id"], "oaiapp_issued")
        self.assertNotIn("agent_name_hint", second_authorize)

    def test_sign_in_failures_do_not_store_anything(self) -> None:
        def denied(url: str) -> bool:
            query = self._query(url)
            target = urllib.parse.urlparse(query["redirect_uri"])

            def callback() -> None:
                conn = http.client.HTTPConnection(target.hostname, target.port, timeout=5)
                conn.request("GET", f"{target.path}?error=access_denied&state={query['state']}")
                conn.getresponse().read()
                conn.close()

            threading.Thread(target=callback).start()
            return True

        with self.assertRaises(chatgpt_auth.ChatGPTSignInCancelled):
            chatgpt_auth.sign_in("p2", open_url=denied, timeout=10)
        with self.assertRaisesRegex(chatgpt_auth.ChatGPTAuthError, "browser"):
            chatgpt_auth.sign_in("p2", open_url=lambda _url: False, timeout=10)
        self.assertEqual(self.keychain.blobs, {})

    def _seed_session(self, profile: str, *, expires_in: float) -> None:
        session = chatgpt_auth.ChatGPTSession(
            client_id="oaiapp_issued", access_token="old-access", refresh_token="old-refresh",
            expires_at=time.time() + expires_in, scopes=chatgpt_auth.SCOPES.split(), email="student@example.com",
        )
        chatgpt_auth._save_state(profile, chatgpt_auth._StoredState(client_id="oaiapp_issued", last_email=session.email, session=session))

    def test_fresh_session_is_not_refreshed_and_stale_one_rotates_its_tokens(self) -> None:
        fake = _FakeTokenServer(lambda form: (200, {"access_token": "new-access", "refresh_token": "new-refresh", "expires_in": 3600}))
        self.addCleanup(fake.close)
        with patch.object(chatgpt_auth, "TOKEN_URL", fake.url):
            self._seed_session("p3", expires_in=3000)
            self.assertEqual(chatgpt_auth.get_valid_session("p3").access_token, "old-access")
            self.assertEqual(fake.requests, [])

            self._seed_session("p3", expires_in=30)
            refreshed = chatgpt_auth.get_valid_session("p3")
        self.assertEqual((refreshed.access_token, refreshed.refresh_token), ("new-access", "new-refresh"))
        self.assertEqual(fake.requests[0]["grant_type"], "refresh_token")
        self.assertEqual(fake.requests[0]["refresh_token"], "old-refresh")
        self.assertTrue(refreshed.plan_usage_enabled, "scopes survive a refresh that doesn't repeat them")
        stored = chatgpt_auth._load_state("p3").session
        self.assertEqual(stored.refresh_token, "new-refresh")

    def test_disconnected_grant_clears_tokens_but_a_network_error_does_not(self) -> None:
        fake = _FakeTokenServer(lambda form: (400, {"error": "invalid_grant"}))
        self.addCleanup(fake.close)
        self._seed_session("p4", expires_in=1)
        with patch.object(chatgpt_auth, "TOKEN_URL", fake.url):
            with self.assertRaisesRegex(chatgpt_auth.ChatGPTAuthError, "sign in again"):
                chatgpt_auth.get_valid_session("p4")
        state = chatgpt_auth._load_state("p4")
        self.assertIsNone(state.session)
        self.assertEqual(state.client_id, "oaiapp_issued")

        self._seed_session("p5", expires_in=1)
        dead_url = fake.url
        fake.close()
        with patch.object(chatgpt_auth, "TOKEN_URL", dead_url):
            with self.assertRaises(chatgpt_auth.ChatGPTNetworkError):
                chatgpt_auth.get_valid_session("p5")
        self.assertIsNotNone(chatgpt_auth._load_state("p5").session)


def _events(*events):
    return iter(events)


class _FakeResponses:
    def __init__(self, outcome) -> None:
        self.calls: list[dict] = []
        self._outcome = outcome

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if isinstance(self._outcome, Exception):
            raise self._outcome
        return self._outcome


class _FakeSDK:
    def __init__(self, outcome) -> None:
        self.responses = _FakeResponses(outcome)
        self.closed = False

    def close(self) -> None:
        self.closed = True


class AIBackendTests(unittest.TestCase):
    CALL = dict(system="be precise", user_content=[{"type": "input_text", "text": "hi"}], text_format={"format": {"type": "json_schema"}}, timeout=5)

    def test_model_choice_prefers_the_requested_model_then_a_mini(self) -> None:
        self.assertEqual(ai_backend.choose_plan_model(["a", "b-mini", "c"], "c"), "c")
        self.assertEqual(ai_backend.choose_plan_model(["big", "x-mini"], "gpt-missing"), "x-mini")
        self.assertEqual(ai_backend.choose_plan_model(["big", "bigger"], "gpt-missing"), "big")
        with self.assertRaises(ai_backend.AIUnavailableError):
            ai_backend.choose_plan_model([], "anything")

    def test_api_key_requests_keep_their_original_shape(self) -> None:
        sdk = _FakeSDK(SimpleNamespace(output_text="{}"))
        client = ai_backend.AIClient(sdk, plan_usage=False, model="gpt-x")
        self.assertEqual(client.generate_text(**self.CALL), "{}")
        call = sdk.responses.calls[0]
        self.assertEqual([item["role"] for item in call["input"]], ["system", "user"])
        self.assertNotIn("stream", call)
        self.assertNotIn("store", call)

    def test_plan_requests_use_instructions_stream_and_do_not_store(self) -> None:
        events = _events(
            SimpleNamespace(type="response.created"),
            SimpleNamespace(type="response.output_text.delta", delta='{"a":'),
            SimpleNamespace(type="response.output_text.delta", delta=" 1}"),
            SimpleNamespace(type="response.completed", response=SimpleNamespace(output_text="")),
        )
        sdk = _FakeSDK(events)
        client = ai_backend.AIClient(sdk, plan_usage=True, model="gpt-x")
        self.assertEqual(client.generate_text(**self.CALL), '{"a": 1}')
        call = sdk.responses.calls[0]
        self.assertIs(call["stream"], True)
        self.assertIs(call["store"], False)
        self.assertEqual(call["instructions"], "be precise")
        self.assertEqual([item["role"] for item in call["input"]], ["user"])
        for unsupported in ("temperature", "max_output_tokens", "previous_response_id"):
            self.assertNotIn(unsupported, call)

    def test_plan_failures_become_actionable_messages(self) -> None:
        failed = SimpleNamespace(
            type="response.failed",
            response=SimpleNamespace(error=SimpleNamespace(code="subscription_sharing_usage_limit_exceeded", message="x", param=None)),
        )
        with self.assertRaisesRegex(ai_backend.AIUnavailableError, "Usage limit reached"):
            ai_backend.AIClient(_FakeSDK(_events(failed)), plan_usage=True, model="m").generate_text(**self.CALL)

        with self.assertRaisesRegex(ai_backend.AIUnavailableError, "interrupted"):
            ai_backend.AIClient(_FakeSDK(_events(SimpleNamespace(type="response.output_text.delta", delta="x"))), plan_usage=True, model="m").generate_text(**self.CALL)

        http_error = type("E", (Exception,), {})("boom")
        http_error.status_code = 403
        http_error.body = {"error": {"code": "subscription_sharing_user_not_eligible", "message": "nope"}}
        with self.assertRaisesRegex(ai_backend.AIUnavailableError, "Plus or Pro"):
            ai_backend.AIClient(_FakeSDK(http_error), plan_usage=True, model="m").generate_text(**self.CALL)

        unsupported = type("E", (Exception,), {})("boom")
        unsupported.status_code = 400
        unsupported.body = {"error": {"code": "subscription_sharing_unsupported_capability", "message": "no pdfs", "param": "input"}}
        with self.assertRaisesRegex(ai_backend.AIUnavailableError, r"'input'.*no pdfs"):
            ai_backend.AIClient(_FakeSDK(unsupported), plan_usage=True, model="m").generate_text(**self.CALL)

    def test_open_client_prefers_chatgpt_then_api_key_then_explains_what_to_do(self) -> None:
        created: list[dict] = []

        class FakeOpenAI:
            def __init__(self, **kwargs) -> None:
                created.append(kwargs)

        session = chatgpt_auth.ChatGPTSession(
            client_id="c", access_token="tok", refresh_token="r", expires_at=time.time() + 3600, scopes=["openid", chatgpt_auth.PLAN_USAGE_SCOPE]
        )
        auth = ai_backend.AIAuth(profile_id="p", api_key="sk-legacy")
        with patch.object(ai_backend, "OpenAIClient", FakeOpenAI):
            with patch.object(chatgpt_auth, "get_valid_session", return_value=session), patch.object(chatgpt_auth, "list_models", return_value=["other", "m-mini"]):
                client = ai_backend.open_ai_client(auth, model="gpt-missing", timeout=9)
            self.assertEqual((client.plan_usage, client.model), (True, "m-mini"))
            self.assertEqual((created[-1]["api_key"], created[-1]["base_url"], created[-1]["max_retries"]), ("tok", chatgpt_auth.API_BASE_URL, 0))

            no_plan = chatgpt_auth.ChatGPTSession(client_id="c", access_token="t", refresh_token="r", expires_at=time.time() + 3600, scopes=["openid"])
            with patch.object(chatgpt_auth, "get_valid_session", return_value=no_plan):
                with self.assertRaisesRegex(ai_backend.AIUnavailableError, "plan"):
                    ai_backend.open_ai_client(auth, model="m", timeout=9)

            with patch.object(chatgpt_auth, "get_valid_session", side_effect=chatgpt_auth.ChatGPTAuthError("Sign in again please.")):
                with self.assertRaisesRegex(ai_backend.AIUnavailableError, "Sign in again please"):
                    ai_backend.open_ai_client(auth, model="m", timeout=9)

            with patch.object(chatgpt_auth, "get_valid_session", return_value=None):
                legacy = ai_backend.open_ai_client(auth, model="gpt-x", timeout=9)
                self.assertEqual((legacy.plan_usage, legacy.model, created[-1]["api_key"]), (False, "gpt-x", "sk-legacy"))
                with self.assertRaisesRegex(ai_backend.AIUnavailableError, "Continue with ChatGPT"):
                    ai_backend.open_ai_client(ai_backend.AIAuth(profile_id="p"), model="m", timeout=9)

    def test_extractors_report_the_sign_in_hint_when_nothing_is_configured(self) -> None:
        from syllabus_import import OpenAISyllabusExtractor, SyllabusExtractionError
        from transcript_import import OpenAITranscriptExtractor, TranscriptExtractionError

        with patch.object(chatgpt_auth, "get_valid_session", return_value=None):
            with self.assertRaisesRegex(SyllabusExtractionError, "Continue with ChatGPT"):
                OpenAISyllabusExtractor(ai_backend.AIAuth(profile_id="p")).extract("/nonexistent.pdf")
            with self.assertRaisesRegex(TranscriptExtractionError, "Continue with ChatGPT"):
                OpenAITranscriptExtractor(ai_backend.AIAuth(profile_id="p")).extract("/nonexistent.pdf")


class RegistryResilienceTests(unittest.TestCase):
    def test_registry_recovers_from_last_valid_backup(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            primary = root / "users.json"
            backup = root / "users.json.bak"
            first = {"version": 1, "current_user_id": "one", "profiles": [{"id": "one", "name": "One"}]}
            second = {"version": 1, "current_user_id": "two", "profiles": [{"id": "two", "name": "Two"}]}

            with (
                patch.object(paths, "APP_SUPPORT_DIR", root),
                patch.object(user_profiles, "REGISTRY_PATH", primary),
                patch.object(user_profiles, "REGISTRY_BACKUP_PATH", backup),
            ):
                user_profiles._write_registry(first)
                user_profiles._write_registry(second)
                primary.write_text("{broken", encoding="utf-8")
                recovered = user_profiles._read_registry()

            self.assertEqual(recovered["current_user_id"], "one")
            self.assertEqual(json.loads(backup.read_text(encoding="utf-8"))["current_user_id"], "one")


if __name__ == "__main__":
    unittest.main()
